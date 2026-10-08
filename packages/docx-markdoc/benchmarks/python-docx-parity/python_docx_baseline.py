"""Clean-room python-docx baseline for the greenfield parity benchmark.

Reimplements, from a generic feature list, the kind of per-matter renderer that
safe-docx's `docx-markdoc create` is meant to replace. It covers house style,
title and headings, centred text, quotes, lettered sub-paragraphs, line and
page breaks, legends, signer blocks, next-page sections with unlinked footers,
a PAGE field, bracketed fill-in highlight, and bold and italic text. It also
adds python-docx's native tables, so the baseline is python-docx at its
strongest rather than a strawman. Clause numbers stay literal text because
python-docx has no numbering API.

Usage: python_docx_baseline.py <input.legacy.mdoc> <output.docx>
Run with: uv run --with python-docx==1.2.0 python python_docx_baseline.py ...
"""

import copy
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Times New Roman"
SIG_LINE = "_" * 30
SECTION_RE = re.compile(r'^<!-- section(?:: ([\w-]+))? footer="(.+)" -->$')
SIGNER_RE = re.compile(r"^<signer>(.+?) \| (.+?)</signer>$")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")


def pin_font(run):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(slot), FONT)


def emit(par, text, highlight=True, italic=False, bold=False, size=None):
    """Write inline text: **bold**, *italic*, trailing-backslash breaks, [fill-ins]."""
    lines = text.split("\\\n")
    depth = 0
    for index, line in enumerate(lines):
        for piece in INLINE_RE.split(line):
            if not piece:
                continue
            is_bold = piece.startswith("**")
            is_italic = not is_bold and piece.startswith("*")
            content = piece[2:-2] if is_bold else piece[1:-1] if is_italic else piece
            # Split the piece wherever bracket depth changes highlight state.
            buffer, marked = "", depth > 0
            for char in content:
                opening = char == "["
                if opening:
                    depth += 1
                now = highlight and depth > 0
                if now != marked and buffer:
                    add_run(par, buffer, marked, bold or is_bold, italic or is_italic, size)
                    buffer = ""
                marked = now
                buffer += char
                if char == "]" and depth > 0:
                    depth -= 1
            if buffer:
                add_run(par, buffer, marked, bold or is_bold, italic or is_italic, size)
        if index < len(lines) - 1 and par.runs:
            par.runs[-1].add_break(WD_BREAK.LINE)


def add_run(par, text, marked, bold, italic, size):
    run = par.add_run(text)
    run.bold = bold or None
    run.italic = italic or None
    if size:
        run.font.size = Pt(size)
    if marked:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    pin_font(run)


def page_field(par):
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        run = par.add_run()
        pin_font(run)
        if kind:
            char = OxmlElement("w:fldChar")
            char.set(qn("w:fldCharType"), kind)
            run._element.append(char)
        elif text == " PAGE ":
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._element.append(instr)
        else:
            run.text = text


def build(source, output):
    raw = open(source, encoding="utf-8").read()
    _, _front, body = raw.split("---\n", 2)
    blocks = [b.strip() for b in re.split(r"\n\s*\n", body) if b.strip()]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(slot), FONT)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], margin, Inches(1))
    doc.core_properties.title = "Benchmark resolution"
    doc.core_properties.author = "Benchmark"

    breaks, page_numbers, pending_page_break = [], False, False
    previous, previous_kind = None, None
    for block in blocks:
        if block == "<!-- pagebreak -->":
            pending_page_break = True
            continue
        if block == "<!-- page-numbers -->":
            page_numbers = True
            continue
        section = SECTION_RE.match(block)
        if section:
            breaks.append((previous, section.group(2)))
            continue
        signer = SIGNER_RE.match(block)
        legend = re.match(r"^<legend>(.*)</legend>$", block, re.S)
        if block.startswith("|"):
            rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in block.splitlines()
                    if not re.match(r"^\|[\s:|-]+\|$", line)]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell_par = table.cell(r, c).paragraphs[0]
                    emit(cell_par, value, bold=(r == 0))
            previous, previous_kind = None, "table"
            continue
        par = doc.add_paragraph()
        fmt = par.paragraph_format
        if signer:
            add_run(par, SIG_LINE, False, False, False, None)
            par.runs[-1].add_break(WD_BREAK.LINE)
            emit(par, signer.group(1))
            add_run(par, "\t", False, False, False, None)
            emit(par, signer.group(2))
            fmt.tab_stops.add_tab_stop(Inches(4.25))
            fmt.space_before = Pt(30 if previous_kind == "signer" else 42)
            fmt.keep_together = True
            if previous is not None and previous_kind != "signer":
                previous.paragraph_format.keep_with_next = True
            kind = "signer"
        elif legend:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            emit(par, legend.group(1), highlight=False, italic=True)
            fmt.space_before = Pt(24)
            fmt.keep_together = True
            if previous is not None:
                previous.paragraph_format.keep_with_next = True
            kind = "legend"
        elif block.startswith("# "):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            emit(par, block[2:], bold=True, size=14)
            kind = "title"
        elif block.startswith("## "):
            emit(par, block[3:], bold=True)
            fmt.keep_with_next = True
            fmt.space_before = Pt(10)
            kind = "heading"
        elif re.match(r"^<center>(.*)</center>$", block, re.S):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            emit(par, block[len("<center>"):-len("</center>")])
            kind = "center"
        elif block.startswith("> "):
            fmt.left_indent = fmt.right_indent = Inches(0.5)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            emit(par, block[2:])
            kind = "quote"
        elif re.match(r"^\([a-z]\) ", block):
            fmt.left_indent = Inches(0.5)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            emit(par, block)
            kind = "sub"
        elif block.startswith(("#", "<")):
            raise SystemExit(f"unsupported block: {block[:40]}")
        else:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            emit(par, block)
            kind = "body"
        if pending_page_break and kind not in ("signer", "legend"):
            fmt.page_break_before = True
            pending_page_break = False
        previous, previous_kind = par, kind

    for anchor, _footer in breaks:
        anchor._p.get_or_add_pPr().append(copy.deepcopy(doc.element.body.get_or_add_sectPr()))
    sections = doc.sections
    if page_numbers:
        footer = sections[0].footer
        footer.is_linked_to_previous = False
        fpar = footer.paragraphs[0]
        fpar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fpar.paragraph_format.space_after = Pt(0)
        page_field(fpar)
    for section, (_anchor, text) in zip(list(sections)[1:], breaks):
        section.footer.is_linked_to_previous = False
        fpar = section.footer.paragraphs[0]
        fpar.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fpar.paragraph_format.space_after = Pt(0)
        emit(fpar, text, highlight=False, italic=True)
    doc.save(output)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
